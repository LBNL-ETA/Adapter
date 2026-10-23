import os
import sys
from adapter.to_python import Excel, Db, Db_sqlalchemy
import openpyxl
from openpyxl.workbook.defined_name import DefinedName
import pandas as pd
import pytest
from pandas._testing import assert_frame_equal
import logging

logging.basicConfig(level=logging.DEBUG)

class ExcelTest:
    """to_python.Excel: loading named tables and ranges from test.xlsx."""
    def setup_method(self):
        """Opens test.xlsx with to_python.Excel."""
        self.exl_loader = Excel(os.path.join(os.path.dirname(__file__), "inputs", "test.xlsx"))

    def test_load(self):
        """Excel.load() returns all named tables and ranges, or only the requested
        names of the requested kind.
        """
        dict_of_dfs = self.exl_loader.load(kind="all")
        assert (
            set(dict_of_dfs.keys())
            == {
                "xlsx_table1",
                "xlsx_table2",
                "run_parameters",
                "xlsx_single_col_table",
                "xlsx_named_range1",
                "xlsx_single_col_range",
                "xlsx_single_value_named_range",
            }
        )

        some_tables = self.exl_loader.load(
            data_object_names=["xlsx_table1", "xlsx_table2"], kind="tables"
        )

        assert set(some_tables.keys()) == {"xlsx_table1", "xlsx_table2"}

        some_ranges = self.exl_loader.load(
            data_object_names=["xlsx_named_range1", "xlsx_single_col_range"],
            kind="ranges",
        )

        assert (
            set(some_ranges.keys())
            == {"xlsx_named_range1", "xlsx_single_col_range"}
        )

    def test_get_unnamed_range(self):
        """Excel.get_named_data_object() reads a `sheet!B2:C3` reference as the same
        dataframe as its cells.
        """
        range_name = "test!B2:C3"
        region = self.exl_loader.wb["test"]["B2:C3"]
        df1 = self.exl_loader.convert_data_object_to_df(region, range_name)
        df2 = self.exl_loader.get_named_data_object(range_name)
        assert_frame_equal(df1, df2)

    def test_get_named_range(self):
        """Excel.get_named_data_object() reads a named range as the same dataframe as
        the cells it refers to.
        """
        range_name = "xlsx_named_range1"
        region = self.exl_loader.wb["test"]["J16:K18"]
        df1 = self.exl_loader.convert_data_object_to_df(region, range_name)
        df2 = self.exl_loader.get_named_data_object(range_name)
        assert_frame_equal(df1, df2)

    def test_convert_single_cell_to_df(self):
        """Excel.convert_data_object_to_df() turns a single cell into a 1x1 dataframe
        with the object's name as its column label.
        """
        ws = self.exl_loader.wb["test"]
        cell = ws["B2"]
        df1 = self.exl_loader.convert_data_object_to_df(cell, "Cell")
        df2 = pd.DataFrame([[cell.value]], columns=["Cell"])
        assert_frame_equal(df1, df2)

    def test_convert_multi_cell_to_df(self):
        """Excel.convert_data_object_to_df() uses a range's first row as column labels."""
        ws = self.exl_loader.wb["test"]
        cells = ws["B2:C3"]
        df1 = self.exl_loader.convert_data_object_to_df(cells, "Cells")
        df2 = pd.DataFrame(
            [[cell.value for cell in row] for row in cells[1:]],
            columns=[cell.value for cell in cells[0]],
        )
        assert_frame_equal(df1, df2)

    def test_load_kind_tables_returns_only_tables(self):
        """Excel.load(kind="tables") returns only named tables."""
        assert set(self.exl_loader.load(kind="tables")) == {
            "xlsx_table1",
            "xlsx_table2",
            "run_parameters",
            "xlsx_single_col_table",
        }

    def test_load_kind_ranges_returns_only_ranges(self):
        """Excel.load(kind="ranges") returns only named ranges."""
        assert set(self.exl_loader.load(kind="ranges")) == {
            "xlsx_named_range1",
            "xlsx_single_col_range",
            "xlsx_single_value_named_range",
        }

    def test_load_raises_for_names_not_in_file(self):
        """Excel.load() raises ValueError when a requested name isn't in the workbook."""
        with pytest.raises(ValueError, match="not found in the input file"):
            self.exl_loader.load(data_object_names=["xlsx_table1", "missing_table"])

    def test_load_raises_for_unsupported_kind(self):
        """Excel.load() raises ValueError for a kind other than all, tables or ranges."""
        with pytest.raises(ValueError, match="unsupported value provided for kwarg kind"):
            self.exl_loader.load(kind="charts")

    def test_get_range_with_quoted_sheet_name(self):
        """Excel.get_named_data_object() accepts a quoted sheet name (`'sheet'!B2:C3`)."""
        assert_frame_equal(
            self.exl_loader.get_named_data_object("'test'!B2:C3"),
            self.exl_loader.get_named_data_object("test!B2:C3"),
        )

    def test_get_named_range_with_multiple_regions_raises(self, tmp_path):
        """Excel.get_named_data_object() raises ValueError for a named range that spans
        several regions.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "S"
        ws["A1"], ws["A2"], ws["C1"], ws["C2"] = "h1", 1, "h2", 2
        wb.defined_names["multi"] = DefinedName(
            "multi", attr_text="S!$A$1:$A$2,S!$C$1:$C$2"
        )
        path = str(tmp_path / "multi.xlsx")
        wb.save(path)

        with pytest.raises(ValueError, match="contains more than one region"):
            Excel(path).get_named_data_object("multi")

class TestDb:
    """to_python.Db: loading tables from a sqlite db."""
    def setup_method(self):
        """Creates Db loaders for test.db (with and without already loaded table1)
        and for corrupt.db.
        """
        self.db = Db(file_path=os.path.join(os.path.dirname(__file__), "inputs", "test.db"), pre_existing_keys=['table1'])
        self.good_db = Db(file_path=os.path.join(os.path.dirname(__file__), "inputs", "test.db"))
        self.bad_db = Db(file_path=os.path.join(os.path.dirname(__file__), "inputs", "corrupt.db"))

    def test_load_pre_keys(self):
        """Db.load() raises ValueError for tables already loaded (pre_existing_keys) and
        loads the others.
        """
        # test file load with pre_existing_keys
        with pytest.raises(ValueError):
            self.db.load()
        # test file load with table names and pre_expisting keys
        assert set(self.db.load(table_names=['table2']).keys()) == {'table2'}
        with pytest.raises(ValueError):
            self.db.load(table_names=['table1'])

    def test_load_bad_db(self):
        """Db.load() raises IOError for a db with no tables (corrupt.db)."""
        with pytest.raises(IOError):
            self.bad_db.load()

    def test_load_df(self):
        """Db.load() returns a requested table's columns and rows as a dataframe."""
        expected = pd.DataFrame(
            {"column_A": ["att1", "att2"], "column_B": ["value1", "value2"]}
        )

        assert_frame_equal(self.good_db.load(table_names=['table3'])['table3'], expected)

class TestDbSqlalchemy:
    """to_python.Db_sqlalchemy (Postgres); only the credentials check runs
    without a database server.
    """
    def test_raises_without_secret_file(self, monkeypatch):
        """Db_sqlalchemy() raises ValueError when adapter/Secret.py (the database
        credentials) is missing.
        """
        # a None entry in sys.modules makes `import adapter.Secret` fail,
        # regardless of whether a local Secret.py exists
        monkeypatch.setitem(sys.modules, "adapter.Secret", None)

        with pytest.raises(ValueError, match="Secret.py"):
            Db_sqlalchemy("server.lbl.gov/database")
